"""Download and parse the IGP's official historical seismic catalog."""

import hashlib
import math
import re
from datetime import date, datetime, time, timezone
from io import BytesIO
from zipfile import BadZipFile

import httpx
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from backend.ingestion.base_connector import BaseConnector, ConnectorError


IGP_CATALOG_ENDPOINT = "https://ultimosismo.igp.gob.pe/api/ultimo-sismo/descargar-datos"
IGP_CATALOG_PAGE = "https://ultimosismo.igp.gob.pe/repositorio/datos-sismicos"
MAX_CATALOG_BYTES = 20_000_000
MAX_CATALOG_RECORDS = 100_000
XLSX_CONTENT_TYPES = {
	"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
	"application/octet-stream",
}


class IGPHistoricalCatalogConnector(BaseConnector):
	provider_key = "igp"

	def __init__(
		self,
		endpoint: str = IGP_CATALOG_ENDPOINT,
		timeout_seconds: float = 30,
		transport: httpx.BaseTransport | None = None,
	) -> None:
		self.endpoint = endpoint
		self.timeout_seconds = timeout_seconds
		self.transport = transport
		self.records_skipped = 0

	def fetch_events(self) -> list[dict]:
		self.records_skipped = 0
		workbook_bytes = self._download_catalog()
		try:
			workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
		except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
			raise ConnectorError("IGP historical catalog is not a valid XLSX workbook") from error

		try:
			worksheet = workbook.active
			if worksheet is None:
				raise ConnectorError("IGP historical catalog has no worksheet")
			rows = worksheet.iter_rows(values_only=True)
			headers = self._read_headers(next(rows, None))
			events: list[dict] = []
			for row_number, row in enumerate(rows, start=2):
				if not any(value is not None and str(value).strip() for value in row):
					continue
				if row_number > MAX_CATALOG_RECORDS + 1:
					raise ConnectorError(f"IGP historical catalog exceeds {MAX_CATALOG_RECORDS} records")
				try:
					event = self._normalize_row(row, headers)
					if event is None:
						self.records_skipped += 1
					else:
						events.append(event)
				except (TypeError, ValueError) as error:
					raise ConnectorError(f"Invalid IGP historical catalog row {row_number}: {error}") from error
			if not events:
				raise ConnectorError("IGP historical catalog returned no records")
			return events
		finally:
			workbook.close()

	def _download_catalog(self) -> bytes:
		params = {
			"tipoCatalogo": "Historico",
			"fechaInicio": "1471-01-01",
			"fechaFin": "1959-12-31",
			"minimaMagnitud": "1",
			"maximaMagnitud": "9",
			"minimaProfundidad": "0",
			"maximaProfundidad": "900",
			"latitudNorte": "90",
			"latitudSur": "-90",
			"longitudEste": "180",
			"longitudOeste": "-180",
		}
		chunks: list[bytes] = []
		total_bytes = 0
		try:
			with httpx.Client(timeout=self.timeout_seconds, transport=self.transport) as client:
				with client.stream(
					"GET",
					self.endpoint,
					params=params,
					headers={"Accept": ", ".join(sorted(XLSX_CONTENT_TYPES)), "User-Agent": "SismoPlatform/0.1"},
				) as response:
					response.raise_for_status()
					content_type = response.headers.get("content-type", "").split(";", maxsplit=1)[0].strip().lower()
					if content_type and content_type not in XLSX_CONTENT_TYPES:
						raise ConnectorError(f"IGP catalog returned an unexpected content type: {content_type}")
					for chunk in response.iter_bytes():
						total_bytes += len(chunk)
						if total_bytes > MAX_CATALOG_BYTES:
							raise ConnectorError(f"IGP historical catalog exceeds the {MAX_CATALOG_BYTES}-byte limit")
						chunks.append(chunk)
		except httpx.HTTPError as error:
			raise ConnectorError("IGP historical catalog service is unavailable") from error
		return b"".join(chunks)

	@staticmethod
	def _read_headers(row: tuple[object, ...] | None) -> dict[str, tuple[int, str]]:
		if row is None:
			raise ConnectorError("IGP historical catalog has no header row")
		headers = {
			str(value).strip().casefold(): (index, str(value).strip())
			for index, value in enumerate(row)
			if value is not None and str(value).strip()
		}
		for required in ("fecha utc", "hora utc"):
			if required not in headers:
				raise ConnectorError(f"IGP historical catalog is missing the '{required}' column")
		for required_prefix in ("latitud", "longitud", "profundidad", "magnitud"):
			if not any(name.startswith(required_prefix) for name in headers):
				raise ConnectorError(f"IGP historical catalog is missing a '{required_prefix}' column")
		return headers

	@staticmethod
	def _normalize_row(row: tuple[object, ...], headers: dict[str, tuple[int, str]]) -> dict | None:
		def get_value(name: str) -> object:
			index = headers[name][0]
			return row[index] if index < len(row) else None

		def find_value(prefix: str) -> tuple[object, str]:
			for name, (index, original_name) in headers.items():
				if name.startswith(prefix):
					return (row[index] if index < len(row) else None), original_name
			raise ValueError(f"missing {prefix} column")

		occurred_at = _parse_utc_datetime(get_value("fecha utc"), get_value("hora utc"))
		latitude, _ = find_value("latitud")
		longitude, _ = find_value("longitud")
		depth, _ = find_value("profundidad")
		latitude_value = _parse_number(latitude, "latitude")
		longitude_value = _parse_number(longitude, "longitude")
		depth_value = _parse_number(depth, "depth")
		if latitude_value is None or longitude_value is None or depth_value is None:
			return None

		magnitude_columns = [
			(name, index, original_name)
			for name, (index, original_name) in headers.items()
			if name.startswith("magnitud")
		]
		magnitude_columns.sort(key=lambda item: _magnitude_priority(item[0]))
		catalog_magnitudes: dict[str, float | None] = {}
		selected_magnitude: float | None = None
		selected_type: str | None = None
		for _, index, original_name in magnitude_columns:
			magnitude_type = _magnitude_type(original_name)
			value = _parse_number(row[index] if index < len(row) else None, f"magnitude {magnitude_type}")
			catalog_magnitudes[magnitude_type] = value
			if selected_magnitude is None and value is not None:
				selected_magnitude = value
				selected_type = magnitude_type
		if selected_magnitude is None:
			return None

		identity = f"{occurred_at.isoformat()}|{latitude_value:.6f}|{longitude_value:.6f}"
		external_id = f"igp-historical-{hashlib.sha256(identity.encode('utf-8')).hexdigest()[:48]}"
		return {
			"codigo": external_id,
			"fecha_hora": occurred_at.isoformat(),
			"latitud": latitude_value,
			"longitud": longitude_value,
			"profundidad": depth_value,
			"magnitud": selected_magnitude,
			"tipo_magnitud": selected_type,
			"referencia": "Catálogo histórico IGP (Perú)",
			"source_url": IGP_CATALOG_PAGE,
			"catalog_magnitudes": catalog_magnitudes,
		}


def _parse_utc_datetime(date_value: object, time_value: object) -> datetime:
	if isinstance(date_value, datetime):
		event_date = date_value.date()
	elif isinstance(date_value, date):
		event_date = date_value
	elif isinstance(date_value, str):
		event_date = date.fromisoformat(date_value.strip())
	else:
		raise ValueError("date is missing or invalid")

	if isinstance(time_value, datetime):
		event_time = time_value.time()
	elif isinstance(time_value, time):
		event_time = time_value
	elif isinstance(time_value, str):
		event_time = time.fromisoformat(time_value.strip())
	else:
		raise ValueError("time is missing or invalid")
	return datetime.combine(event_date, event_time).replace(tzinfo=timezone.utc)


def _parse_number(value: object, field_name: str) -> float | None:
	if value is None or (isinstance(value, str) and value.strip() in {"", "-", "—"}):
		return None
	try:
		number = float(value)
	except (TypeError, ValueError) as error:
		raise ValueError(f"{field_name} is not numeric") from error
	if not math.isfinite(number):
		raise ValueError(f"{field_name} is not finite")
	return number


def _magnitude_type(header: str) -> str:
	match = re.search(r"\(([^)]+)\)", header)
	return match.group(1).strip().upper() if match else "M"


def _magnitude_priority(header: str) -> int:
	magnitude_type = _magnitude_type(header).casefold()
	return {"mw": 0, "ms": 1, "mb": 2, "m": 3}.get(magnitude_type, 4)
